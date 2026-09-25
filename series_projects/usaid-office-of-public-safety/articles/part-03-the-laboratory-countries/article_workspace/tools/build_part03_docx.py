from __future__ import annotations

import csv
import re
import zipfile
from pathlib import Path
from xml.etree import ElementTree as ET

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(r"C:\Users\domin\PycharmProjects\White-Rabbit-Article-Research-Team")
WORKSPACE = ROOT / "series_projects" / "usaid-office-of-public-safety" / "articles" / "part-03-the-laboratory-countries" / "article_workspace"
MARKDOWN_PATH = WORKSPACE / "PART_03_FINAL.md"
LEDGER_PATH = WORKSPACE / "PART_03_SOURCE_LINKS.csv"
OUTPUT_PATH = WORKSPACE / "PART_03_FINAL.docx"

FONT_NAME = "Arial"
BLACK = "000000"
BLUE = "1F4E79"
MID_GRAY = "666666"
LIGHT_GRAY = "D9D9D9"
PALE_GRAY = "F2F2F2"


def set_run_font(run, size=10.5, bold=None, italic=None, color=BLACK, underline=None):
    run.font.name = FONT_NAME
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:ascii"), FONT_NAME)
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:hAnsi"), FONT_NAME)
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if underline is not None:
        run.underline = underline


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=90, start=110, bottom=90, end=110):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, color=LIGHT_GRAY, size="4"):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = borders.find(qn(f"w:{edge}"))
        if tag is None:
            tag = OxmlElement(f"w:{edge}")
            borders.append(tag)
        tag.set(qn("w:val"), "single")
        tag.set(qn("w:sz"), size)
        tag.set(qn("w:space"), "0")
        tag.set(qn("w:color"), color)


def add_external_hyperlink(paragraph, text, url, *, bold=False, italic=False):
    rel_id = paragraph.part.relate_to(url, RT.HYPERLINK, is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), rel_id)
    run = OxmlElement("w:r")
    run_pr = OxmlElement("w:rPr")
    fonts = OxmlElement("w:rFonts")
    fonts.set(qn("w:ascii"), FONT_NAME)
    fonts.set(qn("w:hAnsi"), FONT_NAME)
    run_pr.append(fonts)
    color = OxmlElement("w:color")
    color.set(qn("w:val"), BLUE)
    run_pr.append(color)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    run_pr.append(underline)
    size = OxmlElement("w:sz")
    size.set(qn("w:val"), "21")
    run_pr.append(size)
    size_cs = OxmlElement("w:szCs")
    size_cs.set(qn("w:val"), "21")
    run_pr.append(size_cs)
    if bold:
        run_pr.append(OxmlElement("w:b"))
    if italic:
        run_pr.append(OxmlElement("w:i"))
    run.append(run_pr)
    text_node = OxmlElement("w:t")
    if text.startswith(" ") or text.endswith(" "):
        text_node.set(qn("xml:space"), "preserve")
    text_node.text = text
    run.append(text_node)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


TOKEN_RE = re.compile(r"\[(?P<label>[^\]]+)\]\((?P<url>https?://[^)]+)\)|\*\*(?P<bold>.+?)\*\*|\*(?P<italic>.+?)\*")


def add_inline_markdown(paragraph, text):
    position = 0
    for match in TOKEN_RE.finditer(text):
        if match.start() > position:
            run = paragraph.add_run(text[position:match.start()])
            set_run_font(run)
        if match.group("url"):
            label = match.group("label")
            is_bold = label.startswith("**") and label.endswith("**")
            is_italic = label.startswith("*") and label.endswith("*") and not is_bold
            if is_bold:
                label = label[2:-2]
            elif is_italic:
                label = label[1:-1]
            add_external_hyperlink(paragraph, label, match.group("url"), bold=is_bold, italic=is_italic)
        elif match.group("bold") is not None:
            run = paragraph.add_run(match.group("bold"))
            set_run_font(run, bold=True)
        elif match.group("italic") is not None:
            run = paragraph.add_run(match.group("italic"))
            set_run_font(run, italic=True)
        position = match.end()
    if position < len(text):
        run = paragraph.add_run(text[position:])
        set_run_font(run)


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    set_run_font(run, size=9, color=MID_GRAY)
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    display = OxmlElement("w:t")
    display.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, separate, display, end])


def configure_styles(doc):
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = FONT_NAME
    normal._element.rPr.rFonts.set(qn("w:ascii"), FONT_NAME)
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), FONT_NAME)
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor(0, 0, 0)
    normal.paragraph_format.space_after = Pt(7)
    normal.paragraph_format.line_spacing = 1.12
    normal.paragraph_format.widow_control = True

    title = styles["Title"]
    title.font.name = FONT_NAME
    title._element.rPr.rFonts.set(qn("w:ascii"), FONT_NAME)
    title._element.rPr.rFonts.set(qn("w:hAnsi"), FONT_NAME)
    title.font.size = Pt(23)
    title.font.bold = True
    title.font.color.rgb = RGBColor(0, 0, 0)
    title.paragraph_format.space_after = Pt(12)
    title.paragraph_format.keep_with_next = True

    heading1 = styles["Heading 1"]
    heading1.font.name = FONT_NAME
    heading1._element.rPr.rFonts.set(qn("w:ascii"), FONT_NAME)
    heading1._element.rPr.rFonts.set(qn("w:hAnsi"), FONT_NAME)
    heading1.font.size = Pt(13)
    heading1.font.bold = True
    heading1.font.color.rgb = RGBColor(0, 0, 0)
    heading1.paragraph_format.space_before = Pt(15)
    heading1.paragraph_format.space_after = Pt(6)
    heading1.paragraph_format.keep_with_next = True
    heading1.paragraph_format.keep_together = True

    heading2 = styles["Heading 2"]
    heading2.font.name = FONT_NAME
    heading2._element.rPr.rFonts.set(qn("w:ascii"), FONT_NAME)
    heading2._element.rPr.rFonts.set(qn("w:hAnsi"), FONT_NAME)
    heading2.font.size = Pt(11.5)
    heading2.font.bold = True
    heading2.font.color.rgb = RGBColor(0, 0, 0)
    heading2.paragraph_format.space_before = Pt(11)
    heading2.paragraph_format.space_after = Pt(5)
    heading2.paragraph_format.keep_with_next = True

    if "Image Note" not in styles:
        image_style = styles.add_style("Image Note", WD_STYLE_TYPE.PARAGRAPH)
    else:
        image_style = styles["Image Note"]
    image_style.font.name = FONT_NAME
    image_style._element.rPr.rFonts.set(qn("w:ascii"), FONT_NAME)
    image_style._element.rPr.rFonts.set(qn("w:hAnsi"), FONT_NAME)
    image_style.font.size = Pt(9)
    image_style.font.italic = True
    image_style.font.color.rgb = RGBColor.from_string(MID_GRAY)
    image_style.paragraph_format.left_indent = Inches(0.25)
    image_style.paragraph_format.right_indent = Inches(0.25)
    image_style.paragraph_format.space_before = Pt(5)
    image_style.paragraph_format.space_after = Pt(8)
    image_style.paragraph_format.keep_together = True


def read_front_matter(lines):
    data = {}
    i = 0
    while i < len(lines) and not lines[i].startswith("# "):
        line = lines[i].strip()
        if line.endswith(":") and i + 1 < len(lines):
            data[line[:-1]] = lines[i + 1].strip()
            i += 2
        else:
            i += 1
    return data, i


def build_document():
    markdown = MARKDOWN_PATH.read_text(encoding="utf-8")
    lines = markdown.splitlines()
    front_matter, title_index = read_front_matter(lines)
    title_text = lines[title_index][2:].strip()

    with LEDGER_PATH.open("r", encoding="utf-8-sig", newline="") as handle:
        ledger = list(csv.DictReader(handle))
    markdown_pairs = [(m.group(1).replace("**", ""), m.group(2)) for m in re.finditer(r"\[([^\]]+)\]\((https?://[^)]+)\)", markdown)]
    ledger_pairs = [(row["phrase"], row["link"]) for row in ledger]
    if markdown_pairs != ledger_pairs:
        raise ValueError("Markdown hyperlink sequence does not exactly match PART_03_SOURCE_LINKS.csv")

    doc = Document()
    configure_styles(doc)
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.72)
    section.bottom_margin = Inches(0.68)
    section.left_margin = Inches(0.82)
    section.right_margin = Inches(0.82)
    section.header_distance = Inches(0.3)
    section.footer_distance = Inches(0.3)

    doc.core_properties.title = title_text
    doc.core_properties.subject = "USAID and CIA investigative series article"
    doc.core_properties.author = "The White Rabbit Report"
    doc.core_properties.keywords = front_matter.get("PRIMARY KEYWORD", "")

    title = doc.add_paragraph(style="Title")
    title.add_run(title_text)

    metadata_rows = [
        ("Meta title", front_matter.get("META TITLE", "")),
        ("Meta description", front_matter.get("META DESCRIPTION", "")),
        ("Slug", front_matter.get("SLUG", "")),
        ("Primary keyword", front_matter.get("PRIMARY KEYWORD", "")),
        ("Secondary keywords", front_matter.get("SECONDARY KEYWORDS", "")),
    ]
    table = doc.add_table(rows=len(metadata_rows), cols=2)
    table.autofit = False
    table.columns[0].width = Inches(1.35)
    table.columns[1].width = Inches(5.45)
    set_table_borders(table)
    for row_index, (label, value) in enumerate(metadata_rows):
        left, right = table.rows[row_index].cells
        left.width = Inches(1.35)
        right.width = Inches(5.45)
        left.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        right.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_shading(left, PALE_GRAY)
        set_cell_margins(left)
        set_cell_margins(right)
        lp = left.paragraphs[0]
        lp.paragraph_format.space_after = Pt(0)
        lr = lp.add_run(label)
        set_run_font(lr, size=9, bold=True)
        rp = right.paragraphs[0]
        rp.paragraph_format.space_after = Pt(0)
        rr = rp.add_run(value)
        set_run_font(rr, size=9)

    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(0)
    spacer.paragraph_format.line_spacing = 0.7

    for line in lines[title_index + 1:]:
        stripped = line.strip()
        if not stripped:
            continue
        if stripped.startswith("## "):
            heading_text = stripped[3:].strip()
            paragraph = doc.add_paragraph(style="Heading 1")
            if heading_text == "FREQUENTLY ASKED QUESTIONS":
                paragraph.paragraph_format.page_break_before = True
            paragraph.add_run(heading_text)
        elif stripped.startswith("### "):
            paragraph = doc.add_paragraph(style="Heading 2")
            paragraph.add_run(stripped[4:].strip())
        elif stripped.startswith("[IMAGE:") and stripped.endswith("]"):
            paragraph = doc.add_paragraph(style="Image Note")
            label = paragraph.add_run("IMAGE NOTE  ")
            set_run_font(label, size=9, bold=True, italic=False, color=MID_GRAY)
            content = stripped[len("[IMAGE:"):-1].strip()
            if " | ALT:" in content:
                description, alt = content.split(" | ALT:", 1)
                body = paragraph.add_run(description.strip())
                set_run_font(body, size=9, italic=True, color=MID_GRAY)
                alt_run = paragraph.add_run(f"\nAlt text: {alt.strip()}")
                set_run_font(alt_run, size=9, italic=True, color=MID_GRAY)
            else:
                body = paragraph.add_run(content)
                set_run_font(body, size=9, italic=True, color=MID_GRAY)
        elif stripped in ("[[SUBSCRIBE]]", "[[SHARE]]"):
            paragraph = doc.add_paragraph()
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.space_before = Pt(5)
            paragraph.paragraph_format.space_after = Pt(5)
            run = paragraph.add_run(stripped.strip("[]"))
            set_run_font(run, size=10, bold=True, color=BLUE)
        elif stripped.startswith("- "):
            paragraph = doc.add_paragraph(style="List Bullet")
            paragraph.paragraph_format.space_after = Pt(3)
            add_inline_markdown(paragraph, stripped[2:])
        else:
            paragraph = doc.add_paragraph()
            add_inline_markdown(paragraph, stripped)

    add_page_number(section.footer.paragraphs[0])
    doc.save(OUTPUT_PATH)
    return ledger_pairs


def verify_docx(expected_pairs):
    with zipfile.ZipFile(OUTPUT_PATH) as archive:
        document_xml = ET.fromstring(archive.read("word/document.xml"))
        relationships_xml = ET.fromstring(archive.read("word/_rels/document.xml.rels"))
    ns = {
        "w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main",
        "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
        "pr": "http://schemas.openxmlformats.org/package/2006/relationships",
    }
    targets = {
        rel.attrib["Id"]: rel.attrib.get("Target", "")
        for rel in relationships_xml.findall("pr:Relationship", ns)
        if rel.attrib.get("Type", "").endswith("/hyperlink")
    }
    actual_pairs = []
    for hyperlink in document_xml.findall(".//w:hyperlink", ns):
        rel_id = hyperlink.attrib.get(f"{{{ns['r']}}}id")
        phrase = "".join(node.text or "" for node in hyperlink.findall(".//w:t", ns))
        actual_pairs.append((phrase, targets.get(rel_id, "")))
    if actual_pairs != expected_pairs:
        raise ValueError(f"DOCX hyperlinks do not match ledger\nExpected: {expected_pairs}\nActual: {actual_pairs}")
    print(f"Created {OUTPUT_PATH}")
    print(f"Verified {len(actual_pairs)} exact DOCX hyperlinks")


if __name__ == "__main__":
    pairs = build_document()
    verify_docx(pairs)
